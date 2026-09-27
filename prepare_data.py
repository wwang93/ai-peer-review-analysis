"""Prepare an authorized Excel workbook; all row-level outputs stay private."""
import argparse
import csv
import os
import re
from pathlib import Path

import openpyxl

BOILERPLATE = [
    'Rather than list these out, summarize them so that I can see more clearly how you understand the feedback you got.',
    'You can describe the suggestions as the areas that the chatbot pinpointed.',
    'describe the suggestions as the areas that the chatbot pinpointed.',
    'Again, summarize here rather than list out each change.',
    'Again, summarize here rather than list each change.',
    'What did the chatbot suggest to you?',
    'What changes did you make to your writing as a result of this feedback?',
]
MISSING = {'', 'not found', 'n/a', 'n/a.', 'na', 'nan'}


def normalize(value):
    text = '' if value is None else str(value)
    return re.sub(r'\s+', ' ', text.replace('_x000D_', ' ').replace('\u00a0', ' ')).strip()


def clean(value):
    text = normalize(value)
    status = 'missing' if text.lower() in MISSING else 'editorial_only' if re.fullmatch(r'\(Implied:.*\)', text, re.I) else 'student_text'
    model = text
    for phrase in BOILERPLATE:
        model = model.replace(phrase, ' ')
    model = re.sub(r'(?<=\.)[DR]8\b', '', model)
    model = normalize(re.sub(r'\(Implied(?: suggestion)?:.*?\)', ' ', model, flags=re.I))
    if model.lower() in MISSING and status != 'editorial_only':
        status = 'missing'
    words = len(re.findall(r"\b[a-zA-Z]+(?:'[a-z]+)?\b", model))
    return dict(original_text=value or '', text=text, model_text=model,
                source_status=status, missing=status != 'student_text',
                word_count=words, short_text=words < 8,
                editorial_marker=bool(re.search(r'\[\.\.\.\]|Implied suggestion|\(implied', text, re.I)))


def prepare(path, out):
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook['raw']
    rows = list(sheet.iter_rows(values_only=True))
    columns = []
    for index, header in enumerate(rows[0]):
        match = re.fullmatch(r'reflection(\d)\.(\d)', str(header))
        if match:
            columns.append((index, *map(int, match.groups())))
    expected = {(r, q) for r in range(1, 5) for q in range(1, 6 if r == 4 else 5)}
    if {(r, q) for _, r, q in columns} != expected or len(columns) != 17:
        raise ValueError('Expected 17 unique reflection columns: four questions in R1-R3 and five in R4.')
    records = []
    seen = set()
    for number, row in enumerate(rows[1:], 2):
        if row[0] is None:
            continue
        sid = f'S{int(row[0]):02d}'
        if sid in seen:
            raise ValueError('Duplicate participant ID in the source worksheet.')
        seen.add(sid)
        for column, rnd, question in columns:
            corpus = 'RQ1' if question == 1 else 'RQ2' if question == 2 else 'RQ3' if question == 3 and rnd < 4 else 'R4_affect' if rnd == 4 and question == 4 else 'context'
            records.append(dict(response_id=f'{sid}_R{rnd}_Q{question}', student_id=sid,
                                round=rnd, question=question, corpus=corpus,
                                **clean(row[column]),
                                source=f'authorized workbook | raw!{openpyxl.utils.get_column_letter(column + 1)}{number}'))
    keep = {r['student_id'] for r in records if not r['missing']}
    records = [r for r in records if r['student_id'] in keep]
    if not records:
        raise ValueError('No substantive participant responses found.')
    (out / 'tables').mkdir(parents=True, exist_ok=True)
    (out / 'models').mkdir(exist_ok=True)
    # Filename retained for compatibility with the recorded 39-student analysis.
    target = out / 'tables' / 'responses_39.csv'
    with target.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    labels = out / 'models' / 'topic_labels.json'
    if not labels.exists():
        labels.write_text((Path(__file__).resolve().parent / 'topic_labels.json').read_text(), encoding='utf-8')
    print(f'Prepared {len(keep)} participants and {len(records)} response cells in the private working directory.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    args = parser.parse_args()
    output = Path(os.environ.get('AI_PEER_REVIEW_WORKDIR', Path(__file__).resolve().parent / 'private_run'))
    prepare(args.workbook, output)
